from io import BytesIO

import pytest

from job_agent import config, embeddings, resume_index
from job_agent.chunking import profile_to_chunks
from job_agent.privacy import redact_pii
from job_agent.resume_loader import ResumeLoadError, load_resume_text
from job_agent.resume_schema import Education, Experience, Project, ResumeProfile
from job_agent.session import delete_session, new_session_id, session_dir, validate_session_id
from job_agent.vector_store import IndexMismatchError, VectorStore


@pytest.fixture(autouse=True)
def tmp_data_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA_DIR", tmp_path / "sessions")


def make_profile() -> ResumeProfile:
    return ResumeProfile(
        skills=["Python", "SQL", "Power BI"],
        experience=[
            Experience(organization="Acme", role="Data Intern", bullets=["Cleaned datasets with pandas"])
        ],
        projects=[
            Project(
                name="Flight Fare Prediction",
                tech_stack=["Python", "scikit-learn", "Streamlit"],
                bullets=["Reduced RMSE by 18% with feature engineering", "Deployed on Streamlit"],
            ),
            Project(
                name="Sales Dashboard",
                tech_stack=["Power BI", "SQL"],
                bullets=["Built an interactive dashboard in Power BI"],
            ),
        ],
        education=[Education(institution="Example Institute", degree="B.Tech CSE")],
        certifications=["Python for Data Science"],
    )


# ---- privacy ---------------------------------------------------------------
def test_redact_pii_masks_contact_details():
    text = (
        "Riya Verma\nriya@example.com | +91 98765 43210 | linkedin.com/in/riya-demo\n"
        "Studied 2022-2026 with CGPA 8.4"
    )
    out = redact_pii(text)
    assert "riya@example.com" not in out and "[EMAIL]" in out
    assert "98765" not in out and "[PHONE]" in out
    assert "linkedin.com/in/riya-demo" not in out and "[URL]" in out
    assert out.splitlines()[0] == "[NAME]"
    assert "2022-2026" in out and "8.4" in out  # years and CGPA must survive


def test_redact_keeps_non_name_first_line():
    assert redact_pii("Data Science Intern - Acme Analytics, Bengaluru").startswith("Data Science")


# ---- loader ----------------------------------------------------------------
def test_load_txt():
    text = load_resume_text(b"Skills: Python, SQL and many other things to pass the length check.", "r.txt")
    assert "Python" in text


def test_load_docx_including_tables():
    from docx import Document

    doc = Document()
    doc.add_paragraph("Skills: Python, SQL and plenty of other words for the length check.")
    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = "Project: Flight Fare Prediction"
    buffer = BytesIO()
    doc.save(buffer)
    text = load_resume_text(buffer.getvalue(), "r.docx")
    assert "Flight Fare" in text and "Python" in text


def test_loader_rejects_bad_type_size_and_empty(monkeypatch):
    with pytest.raises(ResumeLoadError, match="Unsupported"):
        load_resume_text(b"x" * 100, "r.exe")
    monkeypatch.setattr(config, "MAX_UPLOAD_MB", 0.0001)
    with pytest.raises(ResumeLoadError, match="too large"):
        load_resume_text(b"x" * 1000, "r.txt")
    monkeypatch.setattr(config, "MAX_UPLOAD_MB", 5)
    with pytest.raises(ResumeLoadError, match="Almost no text"):
        load_resume_text(b"hi", "r.txt")


# ---- sessions --------------------------------------------------------------
def test_session_id_validation_blocks_path_traversal():
    assert validate_session_id(new_session_id())
    for bad in ("../../etc", "abc", "", "A" * 32, None):
        with pytest.raises(ValueError):
            validate_session_id(bad)


# ---- chunking --------------------------------------------------------------
def test_chunks_are_one_per_bullet_with_context():
    chunks = profile_to_chunks(make_profile())
    ids = {c.id for c in chunks}
    assert {"exp-1-b1", "proj-1-b1", "proj-1-b2", "proj-2-b1", "edu-1", "skills-1", "cert-1"} <= ids
    rmse = next(c for c in chunks if c.id == "proj-1-b1")
    assert "Flight Fare Prediction" in rmse.source and "scikit-learn" in rmse.source
    assert rmse.text == "Reduced RMSE by 18% with feature engineering"  # verbatim


def test_known_skills_include_project_tech():
    assert {"python", "sql", "scikit-learn", "streamlit", "power bi"} <= set(make_profile().known_skills())


# ---- index + retrieval -----------------------------------------------------
def test_index_save_load_and_retrieval_ranking():
    sid = new_session_id()
    resume_index.index_profile(make_profile(), sid)

    results = resume_index.retrieve_evidence(sid, "power bi dashboard", k=2)
    assert results[0][0].id in {"proj-2-b1", "skills-1"}
    assert results[0][1] >= results[1][1]  # sorted by score

    results = resume_index.retrieve_evidence(sid, "rmse feature engineering", k=1)
    assert results[0][0].id == "proj-1-b1"

    assert resume_index.load_profile(sid).projects[0].name == "Flight Fare Prediction"


def test_sessions_are_isolated():
    a, b = new_session_id(), new_session_id()
    resume_index.index_profile(make_profile(), a)
    assert session_dir(a).exists() and not session_dir(b).exists()
    delete_session(a)
    assert not session_dir(a).exists()


def test_embedding_mismatch_is_detected(tmp_path):
    chunks = profile_to_chunks(make_profile())
    vectors = embeddings.embed_passages([c.text for c in chunks])
    VectorStore(chunks, vectors, "fake:model-a").save(tmp_path)
    with pytest.raises(IndexMismatchError):
        VectorStore.load(tmp_path, "fake:model-b")


def test_build_resume_index_sends_only_redacted_text(monkeypatch):
    seen = {}

    def fake_parse(text):
        seen["text"] = text
        return make_profile()

    monkeypatch.setattr(resume_index, "parse_resume", fake_parse)
    raw = (
        b"Riya Verma\nriya@example.com | +91 98765 43210\n"
        b"Skills: Python, SQL. Projects: Flight Fare Prediction with scikit-learn."
    )
    sid = new_session_id()
    resume_index.build_resume_index(raw, "r.txt", sid)
    assert "riya@example.com" not in seen["text"] and "98765" not in seen["text"]
    assert (session_dir(sid) / "vectors.npy").exists()
