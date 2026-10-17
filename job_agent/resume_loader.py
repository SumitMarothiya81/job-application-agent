"""Read a resume file (PDF, DOCX, TXT, MD) and return plain text."""

from io import BytesIO
from pathlib import Path

from job_agent import config

MIN_TEXT_CHARS = 50


class ResumeLoadError(ValueError):
    """Raised with a user-friendly message when a resume can't be read."""


def _read_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(BytesIO(data))
        pages = [(page.extract_text() or "") for page in reader.pages]
    except Exception as exc:  # noqa: BLE001
        raise ResumeLoadError(f"Could not read this PDF: {exc}") from exc
    return "\n".join(pages)


def _read_docx(data: bytes) -> str:
    from docx import Document

    try:
        doc = Document(BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ResumeLoadError(f"Could not read this DOCX: {exc}") from exc
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:  # many resume templates use tables
        for row in table.rows:
            for cell in row.cells:
                lines.extend(p.text for p in cell.paragraphs)
    return "\n".join(lines)


def load_resume_text(data: bytes, filename: str) -> str:
    """Validate the upload and extract its text."""
    ext = Path(filename).suffix.lower()
    if ext not in config.ALLOWED_RESUME_EXTENSIONS:
        allowed = ", ".join(sorted(config.ALLOWED_RESUME_EXTENSIONS))
        raise ResumeLoadError(f"Unsupported file type '{ext}'. Allowed: {allowed}")

    max_bytes = int(config.MAX_UPLOAD_MB * 1024 * 1024)
    if len(data) > max_bytes:
        raise ResumeLoadError(f"File too large (max {config.MAX_UPLOAD_MB:g} MB).")

    if ext == ".pdf":
        text = _read_pdf(data)
    elif ext == ".docx":
        text = _read_docx(data)
    else:
        text = data.decode("utf-8", errors="replace")

    text = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if len(text) < MIN_TEXT_CHARS:
        raise ResumeLoadError(
            "Almost no text found. If this is a scanned PDF/image, OCR is not "
            "supported yet; please upload a text-based PDF or a DOCX."
        )
    return text


def load_resume_file(path: str | Path) -> str:
    path = Path(path)
    return load_resume_text(path.read_bytes(), path.name)
